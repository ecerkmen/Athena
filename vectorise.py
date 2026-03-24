import chromadb
import docx
import fitz
import os
import shutil
from sentence_transformers import SentenceTransformer

chunk_list = []

desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
folder_name = "Athena"
athena_path = os.path.join(desktop_path, folder_name)

if not os.path.exists(athena_path):
    os.mkdir(athena_path)

athena_train_path = os.path.join(athena_path, "athena_train")

if not os.path.exists(athena_train_path):
    os.mkdir(athena_train_path)

while True:

    folder_name = input("Enter the name of the subject folder: ")

    if len(folder_name) == 0:
        break

    name_path = os.path.join(athena_train_path, folder_name)

    if not os.path.exists(name_path):
        os.mkdir(name_path)

    while True:

        path_name = input("Enter pathname for the document: ")

        if len(path_name) == 0:
            break

        shutil.move(path_name, name_path)

for root, dirs, files in os.walk(athena_train_path):
     for file in files:
        if file.endswith(".txt"):
            with open(os.path.join(root, file), "r", encoding="utf-8") as f:
                text = f.read()
        elif file.endswith(".docx"):
            doc = docx.Document(os.path.join(root, file))
            text = "\n".join([para.text for para in doc.paragraphs])
        elif file.endswith(".pdf"):
            doc = fitz.Document(os.path.join(root, file))
            text = "\n".join([page.get_text() for page in doc])

        chunks = text.split("\n\n")

        for chunk in chunks:
            words = chunk.split()
            size = len(words)
            chunk_dict = {"source": root, "text": chunk}

            if size > 200:
                for i in range(0, size, 50):
                    sliced = words[i:i + 100]
                    window_dict = {"source": root, "text": " ".join(sliced)}
                    chunk_list.append(window_dict)
            elif 50 < size < 200:
                    chunk_list.append(chunk_dict)
            elif size < 50:
                continue

model = SentenceTransformer('all-MiniLM-L6-v2')
embeddings = model.encode([chunk["text"] for chunk in chunk_list])

client = chromadb.PersistentClient(path=athena_path)
collection = client.get_or_create_collection(name="athena")
collection.add(
    documents = [chunk["text"] for chunk in chunk_list],
    embeddings = embeddings,
    metadatas = [{"source": chunk["source"]} for chunk in chunk_list],
    ids = [str(i) for i in range(len(chunk_list))]
)